import os
import csv
import pandas as pd
from openpyxl import load_workbook
import config

def get_needed_columns(file_key):
    needed = {config.CASEID_COLUMN}

    for domain in config.VARIABLE_RATIONALE.values():
        for spec in domain["variables"].values():
            if spec.get("column") is not None and spec.get("source_file") == file_key:
                needed.add(spec["column"])
            if spec.get("unit_column") is not None and spec.get("source_file") == file_key:
                needed.add(spec["unit_column"])

    if file_key == "main":
        for spec in config.OUTCOME_VARS.values():
            needed.add(spec["column"])
            if spec.get("unit_column"):
                needed.add(spec["unit_column"])

        needed.update(["REOP30", "READ30", "INTV30"])

    return needed


def convert(file_key):
    source = config.FILE_PATHS[file_key]
    output = os.path.join(
        config.DATA_DIR,
        f"PUF_{file_key.upper()}_2023.tsv"
    )

    print("\n" + "=" * 70)
    print(f"Converting: {os.path.basename(source)}")
    print(f"Output:     {os.path.basename(output)}")
    print("=" * 70)

    needed = get_needed_columns(file_key)

    print(f"Pipeline needs {len(needed)} columns.")

    # Read only the header first.
    header_df = pd.read_excel(source, sheet_name=0, nrows=0)
    headers = list(header_df.columns)

    usable = [c for c in headers if c in needed]
    missing = sorted(needed - set(headers))

    if missing:
        print("WARNING: Missing configured columns:")
        for col in missing:
            print(f"  - {col}")

    print(f"Columns being exported: {len(usable)}")

    # Stream the workbook row-by-row.
    wb = load_workbook(
        source,
        read_only=True,
        data_only=True
    )

    ws = wb[wb.sheetnames[0]]

    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    header_index = {name: i for i, name in enumerate(header_row)}

    selected_indices = [header_index[c] for c in usable]

    with open(output, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f, delimiter="\t")

        writer.writerow(usable)

        row_count = 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            writer.writerow([
                row[i] if i < len(row) else None
                for i in selected_indices
            ])

            row_count += 1

            if row_count % 10000 == 0:
                print(f"  exported {row_count:,} rows...")

    wb.close()

    print(f"Finished: {row_count:,} rows")
    print(f"Created: {output}")


if __name__ == "__main__":
    for key in ("main", "intv", "reop", "read"):
        convert(key)

    print("\nAll four files converted.")
